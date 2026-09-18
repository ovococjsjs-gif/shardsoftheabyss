"""Replay the second pass, assemble Markdown, and create an editable Russian DOCX.
The published 2026-09-16 edition is never written. Run from any directory.
"""
from pathlib import Path
import re,json,hashlib,datetime
from docx import Document
from docx.shared import Mm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ED=Path(__file__).resolve().parents[1]
BASE=ED.parent/'2026-09-16'
ROOT=next(p for p in ED.parents if (p/'.git').is_dir())

def paragraphs(text):return re.split(r'\n\s*\n',text.strip())
def plain(text):return text.replace('**','').replace('*','')
def sha(data):return hashlib.sha256(data).hexdigest()
def elt(tag,attrs=None):
    e=OxmlElement(tag)
    for k,v in (attrs or {}).items():e.set(qn(k),str(v))
    return e

def run_text(p,text):
    bold=italic=False
    for token in re.split(r'(\*\*|\*)',text):
        if token=='**':bold=not bold
        elif token=='*':italic=not italic
        elif token:
            r=p.add_run(token);r.bold=bold;r.italic=italic
    assert not bold and not italic,f'Unbalanced emphasis: {text[:120]}'

def make_docx(chapters):
    doc=Document();style=doc.styles['Normal']
    style.font.name='Times New Roman';style.font.size=Pt(12)
    pf=style.paragraph_format;pf.first_line_indent=Mm(7);pf.line_spacing=1.15
    pf.space_before=Pt(0);pf.space_after=Pt(0);pf.widow_control=True
    pf.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    style.element.get_or_add_rPr().append(elt('w:lang',{'w:val':'ru-RU'}))
    for name in ['Heading 1','Title','Subtitle']:
        st=doc.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor.from_string('252220')
        st.paragraph_format.first_line_indent=Mm(0)
        st.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
    h=doc.styles['Heading 1'];h.font.size=Pt(20);h.font.bold=False
    h.paragraph_format.space_before=Pt(30);h.paragraph_format.space_after=Pt(22)
    h.paragraph_format.page_break_before=True;h.paragraph_format.keep_with_next=True
    doc.styles['Title'].font.size=Pt(30)
    doc.settings.element.insert_element_before(elt('w:autoHyphenation',{'w:val':'true'}),'w:characterSpacingControl')
    doc.settings.element.insert_element_before(elt('w:doNotHyphenateCaps',{'w:val':'true'}),'w:characterSpacingControl')
    doc.settings.element.insert_element_before(elt('w:updateFields',{'w:val':'true'}),'w:compat')
    sec=doc.sections[0]
    sec.page_width=Mm(210);sec.page_height=Mm(297)
    sec.top_margin=sec.bottom_margin=Mm(23);sec.left_margin=Mm(27);sec.right_margin=Mm(22)
    sec.header_distance=sec.footer_distance=Mm(12)
    p=doc.add_paragraph('ОСКОЛКИ БЕЗДНЫ','Title');p.paragraph_format.space_before=Pt(150)
    doc.add_paragraph('Книга первая','Subtitle')
    doc.add_page_break()
    p=doc.add_paragraph('Содержание');p.paragraph_format.first_line_indent=Mm(0)
    p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(14)
    p.runs[0].font.size=Pt(18)
    # Static internal links remain usable without a field update or fixed pagination.
    for n in range(1,36):
        p=doc.add_paragraph();p.paragraph_format.first_line_indent=Mm(0)
        p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT;p.paragraph_format.line_spacing=1
        p.paragraph_format.space_after=Pt(1)
        link=elt('w:hyperlink',{'w:anchor':f'chapter_{n:02}'});r=elt('w:r')
        rp=elt('w:rPr');rp.append(elt('w:color',{'w:val':'403832'}));r.append(rp)
        t=elt('w:t');t.text=f'Глава {n}';r.append(t);link.append(r);p._p.append(link)
    sec=doc.add_section(WD_SECTION_START.NEW_PAGE)
    sec.header.is_linked_to_previous=False;sec.footer.is_linked_to_previous=False
    sec._sectPr.insert_element_before(elt('w:pgNumType',{'w:start':'1'}),'w:cols','w:docGrid')
    p=sec.header.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Mm(0)
    r=p.add_run('ОСКОЛКИ БЕЗДНЫ');r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('777777')
    p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Mm(0)
    r=p.add_run();r.font.size=Pt(9)
    r._r.append(elt('w:fldChar',{'w:fldCharType':'begin'}))
    ins=elt('w:instrText');ins.text=' PAGE ';r._r.append(ins)
    r._r.append(elt('w:fldChar',{'w:fldCharType':'end'}))
    for n,pars in chapters.items():
        p=doc.add_paragraph(f'Глава {n}','Heading 1')
        if n==1:p.paragraph_format.page_break_before=False
        p._p.insert(1 if p._p.pPr is not None else 0,elt('w:bookmarkStart',{'w:id':n,'w:name':f'chapter_{n:02}'}))
        p._p.append(elt('w:bookmarkEnd',{'w:id':n}))
        for text in pars[1:]:
            if not text:continue
            if text.strip() in ('***','* * *'):
                p=doc.add_paragraph('* * *');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.first_line_indent=Mm(0)
                p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(6)
                p.paragraph_format.keep_with_next=True
            else:
                p=doc.add_paragraph();run_text(p,text)
    doc.core_properties.title='Осколки Бездны. Книга первая'
    doc.core_properties.author='';doc.core_properties.last_modified_by=''
    doc.core_properties.subject='Редакторская сборка 17 сентября 2026'
    doc.core_properties.language='ru-RU'
    doc.core_properties.created=doc.core_properties.modified=datetime.datetime(2026,9,17,tzinfo=datetime.timezone.utc)
    doc.save(ED/'shards-of-the-abyss.docx')

def main():
    # Never silently overwrite hand edits to previously generated Markdown/DOCX.
    old_manifest=ED/'manifest.json'
    if old_manifest.exists():
        prev=json.loads(old_manifest.read_text())
        check=dict(prev.get('outputs',{}))
        check.update({e['file']:e['output_sha256'] for e in prev['chapters']})
        for name,digest in check.items():
            path=ED/name
            assert path.exists() and sha(path.read_bytes())==digest, f'Generated file changed outside ledger: {name}'
    base_manifest=json.loads((BASE/'manifest.json').read_text())
    chapters={};entries=[]
    for e in base_manifest['chapters']:
        n=e['chapter'];p=BASE/e['file'];raw=p.read_bytes()
        assert sha(raw)==e['output_sha256'],f'Published source changed: {p}'
        chapters[n]=paragraphs(raw.decode())
        entries.append({'chapter':n,'source':str(p.relative_to(ROOT)),
                        'source_sha256':sha(raw),'file':f'chapters/{n:02}.md'})
    log=json.loads((ED/'changes.json').read_text())
    for x in log:
        pars=chapters[x['chapter']];i=x['paragraph']
        assert pars[i]==x['before'],f"Edit mismatch: {x['chapter']}:{i}"
        pars[i]=x['after']
    (ED/'chapters').mkdir(exist_ok=True)
    texts=[]
    # Render inserted multi-paragraph scenes as distinct paragraphs in all formats.
    for entry in entries:
        n=entry['chapter'];text='\n\n'.join(x for x in chapters[n] if x)+'\n'
        chapters[n]=paragraphs(text)
        (ED/entry['file']).write_text(text);texts.append(text)
        entry['output_sha256']=sha(text.encode())
    full='# Осколки Бездны\n\nКнига первая\n\n'+'\n'.join(texts)
    (ED/'book-edited.md').write_text(full)
    make_docx(chapters)
    manifest={'edition':'2026-09-17','status':'second-pass-delivered','base_commit':'690c7bb07b41ef6077f83cc89e0e69bd74911dc7',
              'chapters':entries,'operations':len(log),'outputs':{name:sha((ED/name).read_bytes()) for name in ['book-edited.md','shards-of-the-abyss.docx']}}
    (ED/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
    print(f'Built 35 full chapters, Markdown and DOCX; {len(log)} second-pass operations.')
if __name__=='__main__':main()
